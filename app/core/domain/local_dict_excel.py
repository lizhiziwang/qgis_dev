"""本地库字典表的 Excel 导入 / 导出（.xls / .xlsx）。

格式与「字典管理」导出的业务字典一致：

``字典域 | 代码 | 名称 | 上级代码(仅多级字典需填写)``

- 导出用 xlwt 写 ``.xls``（QGIS 运行时自带）；
- 导入 ``.xls`` 用 xlrd、``.xlsx`` 用 openpyxl；表头可出现在前 10 行内；
- 写回本地字典表时按 ``(domain, code)`` upsert，可先清空再导入。
"""

from __future__ import annotations

import os
import re
from collections.abc import Sequence
from dataclasses import dataclass

from app.core.util.local_db_util import local_db_util
from app.infra.logging import get_logger

logger = get_logger(__name__)

EXPORT_HEADERS: tuple[str, ...] = (
    "字典域",
    "代码",
    "名称",
    "上级代码(仅多级字典需填写)",
)
IMPORT_FILTER = "Excel 字典 (*.xls *.xlsx);;所有文件 (*)"
EXPORT_FILTER = "Excel 97-2003 工作簿 (*.xls);;所有文件 (*)"

_HEADER_ALIASES = {
    "字典域": "domain",
    "字典表": "domain",
    "字典名": "domain",
    "域": "domain",
    "代码": "code",
    "编码": "code",
    "代码值": "code",
    "名称": "name",
    "名称值": "name",
    "显示名": "name",
    "上级代码": "parent_code",
    "上级编码": "parent_code",
    "父代码": "parent_code",
    "父级代码": "parent_code",
    "上级": "parent_code",
}
_PAREN_RE = re.compile(r"[（(][^）)]*[）)]")


class LocalDictExcelError(RuntimeError):
    """本地字典 Excel 读写失败。消息可直接展示。"""


@dataclass
class LocalDictImportReport:
    rows_total: int = 0
    inserted: int = 0
    updated: int = 0
    skipped: int = 0


def _quote(name: str) -> str:
    return '"' + str(name).replace('"', '""') + '"'


def _normalize_header(raw: object) -> str:
    text = _PAREN_RE.sub("", str(raw or ""))
    return text.replace(" ", "").replace("　", "").replace("*", "").strip()


def _header_columns(values: Sequence[object]) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, value in enumerate(values):
        key = _HEADER_ALIASES.get(_normalize_header(value))
        if key and key not in columns:
            columns[key] = index
    return columns


def _resolve_header(rows: Sequence[Sequence[str]]) -> tuple[dict[str, int], int]:
    """返回 ``(列映射, 表头行号)``；表头在前 10 行内。"""
    for index, row in enumerate(rows[:10]):
        columns = _header_columns(row)
        if {"domain", "code", "name"} <= set(columns):
            return columns, index
    raise LocalDictExcelError("认不出表头：文件里需要有「字典域 / 代码 / 名称」三列")


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("　", " ").strip()


def _read_xls(path: str) -> list[list[str]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover
        raise LocalDictExcelError("本机缺少 xlrd，无法读取 .xls") from exc
    try:
        book = xlrd.open_workbook(path, on_demand=True)
        sheet = book.sheet_by_index(0)
        rows = [
            [_cell(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]
        book.release_resources()
        return rows
    except Exception as exc:
        raise LocalDictExcelError(f"读 Excel 失败：{path}\n{exc}") from exc


def _read_xlsx(path: str) -> list[list[str]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise LocalDictExcelError("本机缺少 openpyxl，无法读取 .xlsx") from exc
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = [
            [_cell(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        workbook.close()
        return rows
    except Exception as exc:
        raise LocalDictExcelError(f"读 Excel 失败：{path}\n{exc}") from exc


def _read_rows(path: str) -> list[list[str]]:
    if not os.path.isfile(path):
        raise LocalDictExcelError(f"文件不存在：{path}")
    if os.path.splitext(path)[1].lower() == ".xls":
        return _read_xls(path)
    return _read_xlsx(path)


# --------------------------------------------------------------------- 导出


def export_local_dict_xls(path: str, db_path: str, table: str) -> tuple[str, int]:
    """把本地字典表导出成 ``.xls``，返回 ``(路径, 行数)``。"""
    db = local_db_util(db_path)
    try:
        columns = {
            str(row[1]) for row in db.cursor.execute(f"PRAGMA table_info({_quote(table)})")
        }
        if not {"domain", "code", "name"} <= columns:
            raise LocalDictExcelError("这张表不是字典表结构（缺 domain / code / name 列）")
        parent = "COALESCE(parent_code, '')" if "parent_code" in columns else "''"
        order = "sort_order, rowid" if "sort_order" in columns else "rowid"
        rows = db.cursor.execute(
            f"SELECT domain, code, name, {parent} FROM {_quote(table)} "
            f"ORDER BY domain, {order}"
        ).fetchall()
    except Exception as exc:
        raise LocalDictExcelError(f"读字典表失败：{exc}") from exc
    finally:
        db.conn.close()
    try:
        import xlwt
    except ImportError as exc:  # pragma: no cover
        raise LocalDictExcelError("本机缺少 xlwt，无法导出 .xls") from exc
    target = str(path).strip()
    if not target.lower().endswith(".xls"):
        target = os.path.splitext(target)[0] + ".xls"
    try:
        workbook = xlwt.Workbook(encoding="utf-8")
        sheet = workbook.add_sheet("字典表")
        for column, header in enumerate(EXPORT_HEADERS):
            sheet.write(0, column, header)
        for row_index, row in enumerate(rows, start=1):
            for column, value in enumerate(row):
                sheet.write(row_index, column, str(value or ""))
        workbook.save(target)
    except OSError as exc:
        raise LocalDictExcelError(
            f"写 Excel 失败：{target}\n目标文件可能正被 Excel / WPS 打开，请关闭后重试。\n{exc}"
        ) from exc
    return target, len(rows)


# --------------------------------------------------------------------- 导入


def import_local_dict_xls(
    path: str,
    db_path: str,
    table: str,
    *,
    replace: bool = False,
) -> LocalDictImportReport:
    """把 Excel 字典写回本地字典表（按 ``(domain, code)`` upsert）。"""
    rows = _read_rows(path)
    columns, header_index = _resolve_header(rows)
    report = LocalDictImportReport()
    db = local_db_util(db_path)
    try:
        table_columns = {
            str(row[1]) for row in db.cursor.execute(f"PRAGMA table_info({_quote(table)})")
        }
        if not {"domain", "code", "name"} <= table_columns:
            raise LocalDictExcelError("目标表不是字典表结构（缺 domain / code / name 列）")
        has_parent = "parent_code" in table_columns
        has_sort = "sort_order" in table_columns
        if replace:
            db.cursor.execute(f"DELETE FROM {_quote(table)}")
        index_name = f"ux_{table}_domain_code".replace('"', "")
        db.cursor.execute(
            f'CREATE UNIQUE INDEX IF NOT EXISTS "{index_name}" '
            f'ON {_quote(table)} (domain, code)'
        )
        if has_parent:
            parent_index = f"ux_{table}_domain_parent".replace('"', "")
            db.cursor.execute(
                f'CREATE INDEX IF NOT EXISTS "{parent_index}" '
                f'ON {_quote(table)} (domain, parent_code)'
            )
        existing = {
            str(row[0])
            for row in db.cursor.execute(f"SELECT code FROM {_quote(table)}")
        }
        next_order: dict[tuple[str, str], int] = {}
        domains: set[str] = set()
        for row in rows[header_index + 1:]:
            def cell(key: str, row=row) -> str:
                index = columns.get(key, -1)
                return row[index].strip() if 0 <= index < len(row) else ""

            domain = cell("domain")
            code = cell("code")
            name = cell("name")
            report.rows_total += 1
            if not (domain and code and name):
                report.skipped += 1
                continue
            domains.add(domain)
            parent = cell("parent_code")
            order_key = (domain, parent)
            if order_key not in next_order:
                if has_sort:
                    row_max = db.cursor.execute(
                        f"SELECT COALESCE(MAX(sort_order), -1) FROM {_quote(table)} "
                        "WHERE domain = ? AND COALESCE(parent_code, '') = ?",
                        (domain, parent),
                    ).fetchone()
                    next_order[order_key] = int(row_max[0]) + 1
                else:
                    next_order[order_key] = 0
            order = next_order[order_key]
            next_order[order_key] += 1
            columns_sql = (
                "(domain, code, name, parent_code, sort_order)"
                if has_parent and has_sort
                else "(domain, code, name, parent_code)"
                if has_parent
                else "(domain, code, name)"
            )
            values = [domain, code, name]
            if has_parent:
                values.append(parent)
            if has_sort:
                values.append(order)
            updates = "name = excluded.name, parent_code = excluded.parent_code" if has_parent else "name = excluded.name"
            if has_sort:
                updates += ", sort_order = excluded.sort_order"
            db.cursor.execute(
                f"INSERT INTO {_quote(table)} {columns_sql} "
                f"VALUES ({', '.join('?' for _ in values)}) "
                f"ON CONFLICT(domain, code) DO UPDATE SET {updates}",
                values,
            )
            if code in existing:
                report.updated += 1
            else:
                report.inserted += 1
                existing.add(code)
        with __import__("contextlib").suppress(Exception):
            row = db.cursor.execute(
                "SELECT attr FROM user_relation WHERE table_name = ? AND type = 'dict'",
                (table,),
            ).fetchone()
            import json

            data = {}
            if row and row[0]:
                try:
                    data = json.loads(row[0])
                except (TypeError, ValueError):
                    data = {}
            if isinstance(data, dict):
                all_domains = set(data.get("domains") or []) | domains
                data["domains"] = sorted(str(value) for value in all_domains)
                data.setdefault("schema", "dict_v1")
                data.setdefault("dict_table", table)
                db.cursor.execute(
                    "UPDATE user_relation SET attr = ? "
                    "WHERE table_name = ? AND type = 'dict'",
                    (json.dumps(data, ensure_ascii=False), table),
                )
        db.conn.commit()
    except LocalDictExcelError:
        db.conn.rollback()
        raise
    except Exception as exc:
        db.conn.rollback()
        raise LocalDictExcelError(f"导入字典失败：{exc}") from exc
    finally:
        db.conn.close()
    return report
